"""Compare no skill with the progressive skill.

No-skill runs stay in model-runs/<key>/baseline.
With-skill runs stay in model-runs/<key>-v2/treatment.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
sys.path.insert(0, str(ROOT / "scripts"))

from experiment import load_experiment  # noqa: E402
from model_sweep_results import gold_mcp_calls  # noqa: E402
from model_sweep_xlsx import estimated_cost, mcp_call_count, stamp_workbook  # noqa: E402
from report import passes, percentile  # noqa: E402
from results_md import table  # noqa: E402
from score import read_jsonl, score_item  # noqa: E402

EXPERIMENTS = ("linear", "slack")
V2_SUFFIX = "v2"
ARMS = (
    ("no skill", "baseline", "baseline"),
    ("with skill", "skill", "treatment"),
)
HEADER = PatternFill("solid", fgColor="1F2937")
SUBHEADER = PatternFill("solid", fgColor="E5E7EB")
GOOD = PatternFill("solid", fgColor="D1FAE5")
BAD = PatternFill("solid", fgColor="FEE2E2")
FLAT = PatternFill("solid", fgColor="F3F4F6")


def load_runs(experiment_name: str, key: str) -> dict[tuple[str, str, int], dict[str, Any]]:
    experiment = load_experiment(experiment_name)
    runs: dict[tuple[str, str, int], dict[str, Any]] = {}
    root = experiment.directory / "model-runs" / key
    for path in root.glob("*/*/*.json"):
        run = json.loads(path.read_text())
        runs[(run["arm"], run["question_id"], run["attempt"])] = run
    return runs


def require_v2(experiment_name: str, questions: dict[str, Any], runs: dict) -> None:
    expected = len(questions) * 3
    treatment = [run for (arm, _, _), run in runs.items() if arm == "treatment"]
    problems: list[str] = []
    if len(treatment) != expected:
        problems.append(f"{len(treatment)} treatment trials, expected {expected}")
    if any(arm != "treatment" for arm, _, _ in runs):
        problems.append("v2 run-set contains a non-treatment arm")
    for run in treatment:
        question = questions[run["question_id"]]
        if passes(score_item(question, run)) is None:
            problems.append(f"{run['question_id']}/{run['attempt']} is unjudged")
        if not run.get("skill_read"):
            problems.append(f"{run['question_id']}/{run['attempt']} did not read SKILL.md")
        if run.get("isolation_violation"):
            problems.append(f"{run['question_id']}/{run['attempt']} isolation violation")
        if (run.get("timing") or {}).get("adapter_elapsed_seconds") is None:
            problems.append(f"{run['question_id']}/{run['attempt']} has no timing")
        if any(call.get("cache_miss") for call in run["tool_calls"]):
            problems.append(f"{run['question_id']}/{run['attempt']} cache miss")
    if problems:
        preview = "; ".join(problems[:8])
        extra = f" (+{len(problems) - 8} more)" if len(problems) > 8 else ""
        raise SystemExit(f"{experiment_name} v2 is incomplete: {preview}{extra}")


def metrics(runs: list[dict[str, Any]], questions: dict[str, Any]) -> dict[str, float]:
    uncached = cache_write = cache_read = prompt = output = cost = mcp = gold = 0.0
    passed = 0
    latencies: list[float] = []
    by_question: dict[str, list[bool]] = {}
    for run in runs:
        question = questions[run["question_id"]]
        outcome = bool(passes(score_item(question, run)))
        if outcome:
            passed += 1
        by_question.setdefault(run["question_id"], []).append(outcome)
        usage = run.get("usage") or {}
        prompt_tokens = float(usage.get("input_tokens") or 0)
        read_tokens = float(usage.get("cache_read_tokens") or 0)
        write_tokens = float(usage.get("cache_write_tokens") or 0)
        prompt += prompt_tokens
        cache_read += read_tokens
        cache_write += write_tokens
        output += float(usage.get("output_tokens") or 0)
        uncached += max(0.0, prompt_tokens - read_tokens - write_tokens)
        cost += estimated_cost(run)
        latencies.append(float(run["timing"]["adapter_elapsed_seconds"]))
        mcp += mcp_call_count(run)
        gold += gold_mcp_calls(question)
    count = len(runs)
    return {
        "passed": passed,
        "trials": count,
        "questions": len(by_question),
        "pass_at_3": sum(any(outcomes) for outcomes in by_question.values()),
        "strict": sum(all(outcomes) for outcomes in by_question.values()),
        "input": prompt,
        "uncached": uncached,
        "cache_write": cache_write,
        "cache_read": cache_read,
        "output": output,
        "cost": cost,
        "median": percentile(latencies, 0.5) or 0.0,
        "p95": percentile(latencies, 0.95) or 0.0,
        "mcp": mcp / count,
        "gold": gold / count,
    }


def reading(
    summarized: dict[tuple[str, str, str], dict[str, float]],
    models: list[dict[str, Any]],
) -> list[str]:
    def changed(experiment: str, field: str, better) -> list[str]:
        names = []
        for model in models:
            old = summarized[(experiment, model["key"], "no skill")][field]
            new = summarized[(experiment, model["key"], "with skill")][field]
            if better(new, old):
                names.append(model["key"])
        return names

    linear_pass = changed("linear", "passed", lambda new, old: new > old)
    slack_pass_up = changed("slack", "passed", lambda new, old: new > old)
    slack_pass_down = changed("slack", "passed", lambda new, old: new < old)
    cheaper = []
    fewer_calls = []
    for experiment in EXPERIMENTS:
        for model in models:
            key = model["key"]
            label = f"{experiment.capitalize()} {key}"
            old = summarized[(experiment, key, "no skill")]
            new = summarized[(experiment, key, "with skill")]
            if new["cost"] < old["cost"]:
                cheaper.append(label)
            if new["mcp"] < old["mcp"]:
                fewer_calls.append(label)
    return [
        "## Reading the result",
        "",
        f"- Linear trial pass rises from no skill to with skill in "
        f"{len(linear_pass)} of 6 configurations"
        + (f" ({', '.join(linear_pass)})." if linear_pass else "."),
        f"- Slack trial pass rises for {', '.join(slack_pass_up) or 'none'}"
        f" and falls for {', '.join(slack_pass_down) or 'none'}.",
        f"- Estimated list-price cost per trial is lower with the skill than",
        f"  without it in {len(cheaper)} of 12 configurations.",
        f"- Mean MCP calls are lower with the skill than without it in",
        f"  {len(fewer_calls)} of 12 configurations. Skill-file reads are excluded.",
        "- Compare paired medium and high rows per model. Higher effort is not",
        "  uniformly better.",
        "",
    ]


def rate(part: dict[str, float]) -> str:
    return f"{int(part['passed'])}/{int(part['trials'])} ({100 * part['passed'] / part['trials']:.1f}%)"


def chain(*values: str) -> str:
    return " → ".join(values)


def ratio(part: dict[str, float], key: str) -> str:
    return f"{int(part[key])}/{int(part['questions'])}"


def mean_k(part: dict[str, float], key: str, digits: int) -> str:
    return f"{part[key] / part['trials'] / 1000:.{digits}f}k"


def calls(value: float) -> str:
    return f"{value:.2f}"


def attempts_passed(
    runs: dict[tuple[str, str, int], dict[str, Any]],
    questions: dict[str, Any],
    arm: str,
    qid: str,
) -> int:
    return sum(
        bool(passes(score_item(questions[qid], runs[(arm, qid, attempt)])))
        for attempt in (1, 2, 3)
    )


def style_header(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_workbook(
    destination: Path,
    models: list[dict[str, Any]],
    loaded: dict[tuple[str, str, str], dict[tuple[str, str, int], dict[str, Any]]],
    questions: dict[str, dict[str, Any]],
    order: dict[str, list[str]],
    summarized: dict[tuple[str, str, str], dict[str, float]],
) -> None:
    book = Workbook()
    matrix = book.active
    matrix.title = "Per-question matrix"
    matrix.cell(row=1, column=1, value="Server")
    matrix.cell(row=1, column=2, value="Question")
    matrix.cell(row=1, column=3, value="Class")
    column = 4
    for model in models:
        matrix.cell(row=1, column=column, value=model["key"])
        matrix.merge_cells(
            start_row=1, start_column=column, end_row=1, end_column=column + 1
        )
        for offset, label in enumerate(("no skill", "with skill")):
            cell = matrix.cell(row=2, column=column + offset, value=label)
            cell.font = Font(bold=True)
            cell.fill = SUBHEADER
            cell.alignment = Alignment(horizontal="center")
        column += 2
    style_header(matrix, 1, column - 1)
    for index, label in enumerate(("Server", "Question", "Class"), start=1):
        cell = matrix.cell(row=2, column=index, value=label)
        cell.font = Font(bold=True)
        cell.fill = SUBHEADER

    row = 3
    for experiment in EXPERIMENTS:
        for qid in order[experiment]:
            matrix.cell(row=row, column=1, value=experiment)
            matrix.cell(row=row, column=2, value=qid)
            matrix.cell(row=row, column=3, value=questions[experiment][qid]["class"])
            column = 4
            for model in models:
                key = model["key"]
                base = attempts_passed(
                    loaded[(experiment, key, "baseline")], questions[experiment], "baseline", qid
                )
                skill = attempts_passed(
                    loaded[(experiment, key, "skill")], questions[experiment], "treatment", qid
                )
                matrix.cell(row=row, column=column, value=base)
                skill_cell = matrix.cell(row=row, column=column + 1, value=skill)
                skill_cell.fill = GOOD if skill > base else BAD if skill < base else FLAT
                column += 2
            row += 1
    matrix.freeze_panes = "D3"
    matrix.auto_filter.ref = f"A2:{get_column_letter(column - 1)}{row - 1}"
    for index, width in enumerate([12, 42, 18] + [12] * (column - 4), start=1):
        matrix.column_dimensions[get_column_letter(index)].width = width

    summary = book.create_sheet("Summary")
    summary.append(
        [
            "Server",
            "Setting",
            "Arm",
            "Trials passed",
            "Trials",
            "Trial pass rate",
            "Uncached input tokens",
            "Cache write tokens",
            "Cache read tokens",
            "Output tokens",
            "List-price cost",
            "Median latency (s)",
            "p95 latency (s)",
            "Gold mean MCP calls",
            "Mean MCP calls",
        ]
    )
    style_header(summary, 1, 15)
    for experiment in EXPERIMENTS:
        for model in models:
            for label, _version, _arm in ARMS:
                part = summarized[(experiment, model["key"], label)]
                summary.append(
                    [
                        experiment,
                        model["key"],
                        label,
                        int(part["passed"]),
                        int(part["trials"]),
                        part["passed"] / part["trials"],
                        int(round(part["uncached"])),
                        int(round(part["cache_write"])),
                        int(round(part["cache_read"])),
                        int(round(part["output"])),
                        round(part["cost"], 4),
                        round(part["median"], 3),
                        round(part["p95"], 3),
                        round(part["gold"], 3),
                        round(part["mcp"], 3),
                    ]
                )
    for index, width in enumerate(
        [12, 22, 14, 14, 10, 16, 22, 20, 18, 16, 16, 18, 16, 20, 16],
        start=1,
    ):
        summary.column_dimensions[get_column_letter(index)].width = width
    summary.freeze_panes = "A2"
    book.save(destination)
    stamp_workbook(destination)


def main() -> None:
    config = json.loads((ROOT / "model_sweep.json").read_text())
    models = config["models"]
    questions: dict[str, dict[str, Any]] = {}
    order: dict[str, list[str]] = {}
    loaded: dict[tuple[str, str, str], dict[tuple[str, str, int], dict[str, Any]]] = {}
    summarized: dict[tuple[str, str, str], dict[str, float]] = {}

    for experiment in EXPERIMENTS:
        rows = read_jsonl(load_experiment(experiment).questions)
        questions[experiment] = {row["id"]: row for row in rows}
        order[experiment] = [row["id"] for row in rows]
        for model in models:
            key = model["key"]
            baseline = load_runs(experiment, key)
            skill = load_runs(experiment, f"{key}-{V2_SUFFIX}")
            require_v2(f"{experiment}/{key}", questions[experiment], skill)
            loaded[(experiment, key, "baseline")] = baseline
            loaded[(experiment, key, "skill")] = skill
            for label, version, arm in ARMS:
                source = baseline if version == "baseline" else skill
                selected = [run for (run_arm, _, _), run in source.items() if run_arm == arm]
                summarized[(experiment, key, label)] = metrics(
                    selected, questions[experiment]
                )

    lines = [
        "# No skill compared with the progressive skill",
        "",
        "No skill is the saved baseline arm. With skill is the progressive bundle in",
        "`skills/linear/` and `skills/slack/`: a short `SKILL.md` index plus",
        "`references/` files read only for the chosen object. The no-skill arm was",
        "not rerun. The judge, questions, fixtures, and agent prompt match.",
        "",
        "Each arrow below is no skill, then with skill. List price is the mean per",
        "trial. Input tokens are the full prompt count. Latency is",
        "`adapter_elapsed_seconds`. MCP counts exclude skill-file reads. Gold is 0",
        "when the question must confirm or has no primary tool, otherwise the shortest",
        "acceptable sequence.",
        "",
        "The pivotable workbook is [SKILL_COMPARISON.xlsx](SKILL_COMPARISON.xlsx).",
        "",
        "## Headline results",
        "",
    ]
    headline = []
    for model in models:
        key = model["key"]
        for experiment in EXPERIMENTS:
            parts = [
                summarized[(experiment, key, label)]
                for label, _, _ in ARMS
            ]
            headline.append(
                [
                    experiment.capitalize(),
                    key,
                    chain(*(rate(part) for part in parts)),
                    chain(*(ratio(part, "pass_at_3") for part in parts)),
                    chain(*(ratio(part, "strict") for part in parts)),
                    chain(*(mean_k(part, "input", 1) for part in parts)),
                    chain(*(mean_k(part, "output", 2) for part in parts)),
                    chain(*(f"${part['cost'] / part['trials']:.3f}" for part in parts)),
                    chain(*(f"{part['median']:.1f}s" for part in parts)),
                    chain(*(f"{part['p95']:.1f}s" for part in parts)),
                    f"{calls(parts[0]['gold'])}; "
                    + chain(*(calls(part["mcp"]) for part in parts)),
                ]
            )
    lines += table(
        [
            "Server",
            "Model / effort",
            "Trial pass (no skill → with skill)",
            "Pass@3",
            "Strict reliability",
            "Mean input tokens",
            "Mean output tokens",
            "Estimated cost/trial",
            "Median latency",
            "p95 latency",
            "MCP calls (gold; no skill → with skill)",
        ],
        headline,
    )
    lines += reading(summarized, models)

    lines += ["## Per question", ""]
    lines.append(
        "Each cell is attempts passed out of three. With skill is colored in the"
    )
    lines.append("workbook when it differs from no skill.")
    lines.append("")
    for experiment in EXPERIMENTS:
        for model in models:
            key = model["key"]
            lines.append(f"### {experiment} · {key}")
            lines.append("")
            rows = []
            for qid in order[experiment]:
                rows.append(
                    [
                        qid,
                        questions[experiment][qid]["class"],
                        f"{attempts_passed(loaded[(experiment, key, 'baseline')], questions[experiment], 'baseline', qid)}/3",
                        f"{attempts_passed(loaded[(experiment, key, 'skill')], questions[experiment], 'treatment', qid)}/3",
                    ]
                )
            lines.extend(
                table(
                    ["Question", "Class", "No skill", "With skill"],
                    rows,
                )
            )

    results = ROOT / "results" / "comparison"
    results.mkdir(parents=True, exist_ok=True)
    markdown = results / "SKILL_COMPARISON.md"
    workbook = results / "SKILL_COMPARISON.xlsx"
    markdown.write_text("\n".join(lines).rstrip() + "\n")
    write_workbook(workbook, models, loaded, questions, order, summarized)
    print(f"Wrote {markdown}")
    print(f"Wrote {workbook}")


if __name__ == "__main__":
    main()
